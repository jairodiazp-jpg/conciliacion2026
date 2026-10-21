from __future__ import annotations

import base64
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


REQUIRED_FIELDS = ("cuenta", "valor", "fecha", "documento", "sap")
FIELD_ALIASES = {
    "cuenta": ("cuenta", "cuenta contable", "numero de cuenta", "número de cuenta", "cta"),
    "valor": ("valor", "importe", "importe recibo", "monto", "valor total"),
    "fecha": ("fecha", "fecha banco", "fecha contabilizacion", "fecha contabilización", "fecha documento"),
    "documento": ("documento", "numero documento", "número documento", "no documento"),
    "sap": ("sap", "codigo sap", "código sap", "cod sap", "codigo_sap", "código_sap"),
}
PENDING_FIELDS = ("estado", "seguimiento", "observacion", "observaciones", "comentario", "comentarios")
CONCILIATED_MARKERS = ("cruce", "conciliad", "conciliación", "conciliacion")


@dataclass(frozen=True)
class RowRef:
    sheet: str
    row: int


@dataclass
class Record:
    ref: RowRef
    values: dict[str, Any]
    columns: dict[str, int]
    pending: bool
    complete: bool
    reason: str = ""


def _normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(char for char in text if not unicodedata.combining(char)).strip().lower()


def _header_field(value: Any) -> str | None:
    header = _normalize(value)
    compact = re.sub(r"[^a-z0-9]+", " ", header).strip()
    for field, aliases in FIELD_ALIASES.items():
        if compact in {re.sub(r"[^a-z0-9]+", " ", alias).strip() for alias in aliases}:
            return field
    return None


def _parse_decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value.quantize(Decimal("0.01"))
    if isinstance(value, (int, float)):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    text = str(value).strip().replace("$", "").replace(" ", "")
    if not text:
        return None
    if "." in text and "," in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _text(value: Any) -> str:
    return str(value or "").strip()


def _append_once(current: Any, text: str) -> str:
    existing = _text(current)
    if text.lower() in existing.lower():
        return existing
    return f"{existing} | {text}" if existing else text


def _write_traceability(workbook, record: Record, text: str) -> bool:
    preferred_fields = ("comentarios", "comentario", "observaciones", "observacion")
    for field in preferred_fields:
        column = record.columns.get(field)
        if column is not None:
            cell = workbook[record.ref.sheet].cell(record.ref.row, column)
            cell.value = _append_once(cell.value, text)
            return True
    return False


def _is_conciliated(values: dict[str, Any]) -> bool:
    joined = " ".join(_normalize(value) for value in values.values() if value is not None)
    return any(marker in joined for marker in CONCILIATED_MARKERS)


def _find_sections(sheet: Worksheet) -> list[tuple[int, dict[str, int], dict[str, int]]]:
    sections: list[tuple[int, dict[str, int], dict[str, int]]] = []
    for row_number in range(1, sheet.max_row + 1):
        fields: dict[str, int] = {}
        pending: dict[str, int] = {}
        for cell in sheet[row_number]:
            field = _header_field(cell.value)
            if field and field not in fields:
                fields[field] = cell.column
            normalized = _normalize(cell.value)
            if normalized in PENDING_FIELDS:
                pending[normalized] = cell.column
        if fields:
            sections.append((row_number, fields, pending))
    return sections


def _records(workbook, *, include_all: bool, required_values: tuple[str, ...]) -> tuple[list[Record], list[dict[str, Any]]]:
    records: list[Record] = []
    issues: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        sections = _find_sections(sheet)
        for index, (header_row, columns, pending_columns) in enumerate(sections):
            next_header = sections[index + 1][0] if index + 1 < len(sections) else sheet.max_row + 1
            missing = [field for field in REQUIRED_FIELDS if field not in columns]
            if missing:
                issues.append({"sheet": sheet.title, "row": header_row, "type": "columnas_faltantes", "fields": missing})
                continue
            for row_number in range(header_row + 1, next_header):
                raw = {field: sheet.cell(row_number, column).value for field, column in columns.items()}
                if not any(value not in (None, "") for value in raw.values()):
                    continue
                values = {
                    "cuenta": _text(raw["cuenta"]),
                    "valor": _parse_decimal(raw["valor"]),
                    "fecha": _parse_date(raw["fecha"]),
                    "documento": _text(raw["documento"]),
                    "sap": _text(raw["sap"]),
                }
                complete = all(values[field] not in (None, "") for field in required_values)
                status_values = {
                    header: sheet.cell(row_number, column).value
                    for header, column in pending_columns.items()
                }
                pending = not _is_conciliated(status_values)
                if not complete:
                    issues.append({
                        "sheet": sheet.title,
                        "row": row_number,
                        "type": "datos_incompletos",
                        "fields": [field for field in required_values if values[field] in (None, "")],
                    })
                if include_all or pending:
                    records.append(Record(RowRef(sheet.title, row_number), values, columns, pending, complete))
    return records, issues


def _find_header_row(sheet: Worksheet, required: set[str]) -> tuple[int, dict[str, int]] | None:
    for row_number in range(1, min(sheet.max_row, 30) + 1):
        columns: dict[str, int] = {}
        for cell in sheet[row_number]:
            field = _header_field(cell.value)
            if field and field not in columns:
                columns[field] = cell.column
        if required.issubset(columns):
            return row_number, columns
    return None


def _extract_account(value: Any) -> str:
    text = _text(value)
    match = re.search(r"(?:cuenta(?:\s+contable)?|cta)[^0-9]*([0-9][0-9-]*)", _normalize(text))
    if match:
        return re.sub(r"\D", "", match.group(1))
    digits = re.sub(r"\D", "", text)
    return digits


def _query_records(workbook) -> tuple[list[Record], list[dict[str, Any]]]:
    records: list[Record] = []
    issues: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        header = _find_header_row(sheet, {"valor", "fecha"})
        if header is None:
            issues.append({"sheet": sheet.title, "type": "columnas_faltantes", "fields": ["fecha", "valor"]})
            continue
        header_row, columns = header
        bank_col = next(
            (cell.column for cell in sheet[header_row] if _normalize(cell.value) == "banco"),
            None,
        )
        sap_col = next(
            (cell.column for cell in sheet[header_row] if "no documento sap" in _normalize(cell.value)),
            None,
        )
        comments_col = next(
            (
                cell.column
                for cell in sheet[header_row]
                if _normalize(cell.value) in {"comentario", "comentarios", "observacion", "observaciones"}
            ),
            None,
        )
        if comments_col is not None:
            columns["comentarios"] = comments_col
        for row_number in range(header_row + 1, sheet.max_row + 1):
            raw_date = sheet.cell(row_number, columns["fecha"]).value
            raw_value = sheet.cell(row_number, columns["valor"]).value
            if raw_date in (None, "") and raw_value in (None, ""):
                continue
            values = {
                "cuenta": _extract_account(sheet.cell(row_number, bank_col).value) if bank_col else "",
                "valor": _parse_decimal(raw_value),
                "fecha": _parse_date(raw_date),
                "documento": "",
                "sap": _text(sheet.cell(row_number, sap_col).value) if sap_col else "",
            }
            complete = all(values[field] not in (None, "") for field in ("cuenta", "valor", "fecha"))
            if not complete:
                issues.append({
                    "sheet": sheet.title,
                    "row": row_number,
                    "type": "datos_incompletos",
                    "fields": [field for field in ("cuenta", "valor", "fecha") if values[field] in (None, "")],
                })
            records.append(Record(RowRef(sheet.title, row_number), values, columns, True, complete))
    return records, issues


def _memorando_records(workbook) -> tuple[list[Record], list[dict[str, Any]]]:
    records: list[Record] = []
    issues: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        account = ""
        bank_account = ""
        for row_number in range(1, min(sheet.max_row, 12) + 1):
            for cell in sheet[row_number]:
                if "cuenta contable" in _normalize(cell.value):
                    account = _extract_account(cell.value)
                    break
            if account:
                break
        for row_number in range(1, min(sheet.max_row, 12) + 1):
            row_text = " ".join(_text(cell.value) for cell in sheet[row_number])
            if "cuenta corriente" in _normalize(row_text) or "cuenta ahorros" in _normalize(row_text):
                bank_account = _extract_account(row_text)
                if bank_account:
                    break
        sections = _find_sections(sheet)
        for index, (header_row, columns, pending_columns) in enumerate(sections):
            required = {"fecha", "valor", "documento"}
            if not required.issubset(columns):
                continue
            next_header = sections[index + 1][0] if index + 1 < len(sections) else sheet.max_row + 1
            trace_columns = {
                normalized: cell.column
                for cell in sheet[header_row]
                for normalized in (_normalize(cell.value),)
                if normalized in {"comentario", "comentarios", "observacion", "observaciones"}
            }
            for row_number in range(header_row + 1, next_header):
                raw = {field: sheet.cell(row_number, column).value for field, column in columns.items()}
                if raw["fecha"] in (None, "") and raw["valor"] in (None, ""):
                    continue
                values = {
                    "cuenta": bank_account or account or _text(raw.get("cuenta")),
                    "valor": _parse_decimal(raw["valor"]),
                    "fecha": _parse_date(raw["fecha"]),
                    "documento": _text(raw["documento"]),
                    "sap": _text(raw["documento"]),
                }
                if values["fecha"] is None or values["valor"] is None:
                    continue
                status_values = {
                    header: sheet.cell(row_number, column).value
                    for header, column in pending_columns.items()
                }
                pending = not _is_conciliated(status_values)
                records.append(Record(RowRef(sheet.title, row_number), values, trace_columns | columns, pending, True))
    return records, issues


def _record_json(record: Record) -> dict[str, Any]:
    return {
        "sheet": record.ref.sheet,
        "row": record.ref.row,
        "cuenta": record.values["cuenta"],
        "valor": str(record.values["valor"]) if record.values["valor"] is not None else None,
        "fecha": record.values["fecha"].isoformat() if record.values["fecha"] else None,
        "documento": record.values["documento"],
        "sap": record.values["sap"],
    }


def _encode(workbook) -> str:
    output = BytesIO()
    workbook.save(output)
    return base64.b64encode(output.getvalue()).decode("ascii")


def conciliar_query_memorando(query_bytes: bytes, memorando_bytes: bytes) -> dict[str, Any]:
    query_workbook = load_workbook(BytesIO(query_bytes))
    memo_workbook = load_workbook(BytesIO(memorando_bytes))
    query_records, query_issues = _query_records(query_workbook)
    if not query_records or not any(record.values["cuenta"] for record in query_records):
        query_records, query_issues = _records(
            query_workbook,
            include_all=True,
            required_values=("cuenta", "valor", "fecha"),
        )
    memo_records, memo_issues = _memorando_records(memo_workbook)
    if not memo_records or not any(record.values["cuenta"] for record in memo_records):
        memo_records, memo_issues = _records(
            memo_workbook,
            include_all=True,
            required_values=REQUIRED_FIELDS,
        )
    pending_memo = [record for record in memo_records if record.pending]
    used_memo: set[RowRef] = set()
    results: list[dict[str, Any]] = []
    confirmed: list[tuple[Record, Record]] = []
    traceability_written = 0

    for query in query_records:
        if not query.pending:
            results.append({"status": "descartado_ya_conciliado", "query": _record_json(query), "reason": "El registro ya tiene una marca de conciliación."})
            continue
        if not query.complete:
            results.append({"status": "datos_incompletos", "query": _record_json(query), "reason": "Faltan campos obligatorios."})
            continue
        candidates = [
            memo for memo in pending_memo
            if memo.ref not in used_memo
            and all(memo.values[field] not in (None, "") for field in ("cuenta", "valor", "fecha"))
            and memo.values["cuenta"] == query.values["cuenta"]
            and memo.values["valor"] == query.values["valor"]
            and memo.values["fecha"] == query.values["fecha"]
            and (not query.values["documento"] or memo.values["documento"] == query.values["documento"])
        ]
        if len(candidates) > 1:
            results.append({"status": "coincidencia_ambigua", "query": _record_json(query), "candidates": [_record_json(item) for item in candidates], "reason": "Existen varios candidatos exactos."})
            continue
        if len(candidates) == 1:
            memo = candidates[0]
            if not (memo.values["documento"] or query.values["sap"]):
                results.append({"status": "documento_o_sap_no_identificable", "query": _record_json(query), "candidate": _record_json(memo), "reason": "El movimiento contable no tiene documento o código SAP."})
                continue
            used_memo.add(memo.ref)
            confirmed.append((query, memo))
            query_trace = (
                f"Conciliado exacto | Memorando: {memo.ref.sheet} fila {memo.ref.row} | "
                f"Cuenta: {query.values['cuenta']} | Valor: {query.values['valor']} | "
                f"Fecha: {query.values['fecha'].isoformat()} | SAP: {query.values['sap'] or 'No disponible'}"
            )
            memo_trace = (
                f"Conciliado exacto | Query: {query.ref.sheet} fila {query.ref.row} | "
                f"Cuenta: {query.values['cuenta']} | Valor: {query.values['valor']} | "
                f"Fecha: {query.values['fecha'].isoformat()} | SAP Query: {query.values['sap'] or 'No disponible'}"
            )
            if _write_traceability(query_workbook, query, query_trace):
                traceability_written += 1
            if _write_traceability(memo_workbook, memo, memo_trace):
                traceability_written += 1
            source_detail = (
                "documento y SAP provienen del memorando."
                if memo.values["documento"] or memo.values["sap"]
                else "el identificador SAP proviene del Query porque el memorando no tiene una columna Documento diligenciada."
            )
            results.append({
                "status": "conciliado_exacto",
                "query": _record_json(query),
                "memorando": _record_json(memo),
                "reason": f"Cuenta, valor y fecha coinciden exactamente; {source_detail}",
            })
            continue
        same_account_date = [
            memo for memo in pending_memo
            if memo.ref not in used_memo
            and all(memo.values[field] not in (None, "") for field in ("cuenta", "valor", "fecha"))
            and memo.values["cuenta"] == query.values["cuenta"]
            and memo.values["fecha"] == query.values["fecha"]
        ]
        if any(memo.values["valor"] != query.values["valor"] for memo in same_account_date):
            status = "discrepancia_de_valor"
        elif any(memo.values["cuenta"] != query.values["cuenta"] and memo.values["valor"] == query.values["valor"] and memo.values["fecha"] == query.values["fecha"] for memo in pending_memo):
            status = "discrepancia_de_cuenta"
        else:
            status = "sin_coincidencia"
        results.append({"status": status, "query": _record_json(query), "reason": "No existe una correspondencia inequívoca con todos los criterios obligatorios."})

    return {
        "query_file": {"name": "QUERY_CONCILIADO.xlsx", "file": _encode(query_workbook)},
        "memorando_file": {"name": "MEMORANDO_CONCILIADO_QUERY.xlsx", "file": _encode(memo_workbook)},
        "summary": {
            "query_revisados": len(query_records),
            "coincidencias_exactas": len(confirmed),
            "sin_coincidencia": sum(item["status"] == "sin_coincidencia" for item in results),
            "ambiguas": sum(item["status"] == "coincidencia_ambigua" for item in results),
            "incompletos": sum(item["status"] == "datos_incompletos" for item in results),
            "descartados_ya_conciliados": sum(item["status"] == "descartado_ya_conciliado" for item in results),
            "trazabilidades_escritas": traceability_written,
        },
        "issues": query_issues + memo_issues,
        "results": results,
        "warnings": [
            "La trazabilidad se escribió únicamente en columnas existentes de comentarios u observaciones.",
            "No se agregaron hojas, columnas ni se alteraron formatos del memorando.",
        ],
    }
