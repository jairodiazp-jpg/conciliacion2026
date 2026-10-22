import base64
import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from adquirencias_conciliador import AdquirenciasConciliador
from integrador import ProcesadorIntegrado


def _workbook_bytes(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class TestAdquirenciasOutput(unittest.TestCase):
    def test_includes_processed_adquirencias_when_account_is_in_sheet_content(self) -> None:
        ccs = Workbook()
        ccs.remove(ccs.active)
        sheet = ccs.create_sheet("Cruces Bancarios")
        sheet.append(["Cuenta contable", "N° De Aprobación", "Fecha", "Valor"])
        sheet.append(["2490", "166580", "01/05/2025", 12500])

        adquirencias = Workbook()
        sheet = adquirencias.active
        sheet.title = "Adquirencias"
        sheet.append(["Fecha de Transaccion", "Valor Total", "Codigo Autorizacion"])
        sheet.append(["01/05/2025", 12500, "166580"])

        result = ProcesadorIntegrado(
            contable_bytes=_workbook_bytes(ccs),
            adquirencias_bytes=_workbook_bytes(adquirencias),
        ).procesar()

        files = {item["name"]: item["file"] for item in result["files"]}
        self.assertIn("ADQUIRENCIAS_CONCILIADO.xlsx", files)
        self.assertGreater(len(base64.b64decode(files["ADQUIRENCIAS_CONCILIADO.xlsx"])), 0)
        load_workbook(BytesIO(base64.b64decode(files["ADQUIRENCIAS_CONCILIADO.xlsx"])))


if __name__ == "__main__":
    unittest.main()
