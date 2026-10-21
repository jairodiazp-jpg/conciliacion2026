from __future__ import annotations

import base64
import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from pse_conciliador import PINK_FILL, PseConciliador


def _workbook_bytes(title: str, headers: list[str], rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pse(rows: list[list[object]]) -> bytes:
    return _workbook_bytes("Movimientos PSE", ["Fecha", "Valor", "Descripción", "Referencia"], rows)


def _cruces(rows: list[list[object]]) -> bytes:
    return _workbook_bytes("Bancolombia 0531", ["Fecha", "Valor", "Descripción", "Referencia"], rows)


def _states(resultado: dict) -> dict[int, str]:
    return {item["row"]: item["estado_conciliacion"] for item in resultado["dataset"]}


class TestPseExactitud(unittest.TestCase):
    def test_coincidencia_directa_valor_exacto(self) -> None:
        pse = _pse([["2026-04-15", 150000, "Pago PSE", "A1"]])
        cruces = _cruces([["2026-04-15", 150000, "Pago virtual PSE", "B1"]])
        resultado = PseConciliador(pse, cruces).procesar()
        self.assertEqual(_states(resultado)[2], "Conciliado 1:1")

    def test_suma_varios_pse_igual_al_contable(self) -> None:
        pse = _pse(
            [
                ["2026-04-15", 50000, "Pago PSE 1", "P1"],
                ["2026-04-15", 60000, "Pago PSE 2", "P2"],
                ["2026-04-15", 40000, "Pago PSE 3", "P3"],
            ]
        )
        cruces = _cruces([["2026-04-15", 150000, "Pago virtual PSE", "C1"]])
        resultado = PseConciliador(pse, cruces).procesar()
        states = _states(resultado)
        self.assertEqual(states[2], "Conciliado N:1")
        self.assertEqual(states[3], "Conciliado N:1")
        self.assertEqual(states[4], "Conciliado N:1")
        comment = resultado["dataset"][0]["comentario_conciliacion"]
        self.assertIn("Diferencia: 0", comment)
        self.assertIn("Suma calculada: 150,000.00", comment)

        pink = PINK_FILL.start_color.rgb
        pse_wb = load_workbook(BytesIO(base64.b64decode(resultado["file"])))
        for row in range(2, 5):
            self.assertEqual(pse_wb.active.cell(row=row, column=2).fill.start_color.rgb, pink)

    def test_suma_inferior_una_unidad(self) -> None:
        pse = _pse(
            [
                ["2026-04-15", 50000, "Pago PSE 1", "P1"],
                ["2026-04-15", 60000, "Pago PSE 2", "P2"],
                ["2026-04-15", 39999, "Pago PSE 3", "P3"],
            ]
        )
        cruces = _cruces([["2026-04-15", 150000, "Pago virtual PSE", "C1"]])
        resultado = PseConciliador(pse, cruces).procesar()
        self.assertTrue(all(state == "Sin coincidencia" for state in _states(resultado).values()))
        self.assertEqual(resultado["resumen"]["cruzados"], 0)

    def test_suma_superior_una_unidad(self) -> None:
        pse = _pse(
            [
                ["2026-04-15", 50000, "Pago PSE 1", "P1"],
                ["2026-04-15", 60000, "Pago PSE 2", "P2"],
                ["2026-04-15", 40001, "Pago PSE 3", "P3"],
            ]
        )
        cruces = _cruces([["2026-04-15", 150000, "Pago virtual PSE", "C1"]])
        resultado = PseConciliador(pse, cruces).procesar()
        self.assertTrue(all(state == "Sin coincidencia" for state in _states(resultado).values()))

    def test_combinaciones_duplicadas_quedan_pendientes(self) -> None:
        pse = _pse(
            [
                ["2026-04-15", 50000, "Pago PSE 1", "P1"],
                ["2026-04-15", 100000, "Pago PSE 2", "P2"],
                ["2026-04-15", 150000, "Pago PSE 3", "P3"],
            ]
        )
        cruces = _cruces([["2026-04-15", 150000, "Pago virtual PSE", "C1"]])
        resultado = PseConciliador(pse, cruces).procesar()
        self.assertTrue(all(state == "Pendiente de revisión" for state in _states(resultado).values()))
        self.assertEqual(resultado["resumen"]["cruzados"], 0)

    def test_movimiento_pse_ya_utilizado(self) -> None:
        pse = _pse(
            [
                ["2026-04-15", 150000, "Pago PSE", "REF-1"],
                ["2026-04-15", 150000, "Pago PSE 2", "OTRA"],
            ]
        )
        cruces = _cruces([["2026-04-15", 150000, "Pago virtual PSE", "REF-1"]])
        resultado = PseConciliador(pse, cruces).procesar()
        states = _states(resultado)
        self.assertEqual(states[2], "Conciliado por referencia y valor")
        self.assertEqual(states[3], "Sin coincidencia")

    def test_valores_vacios_o_mal_formados_no_confirman(self) -> None:
        pse = _pse(
            [
                ["2026-04-15", None, "Pago PSE", "P1"],
                ["2026-04-15", "abc", "Pago PSE", "P2"],
                ["2026-04-15", 150000, "Pago PSE", "P3"],
            ]
        )
        cruces = _cruces([["2026-04-15", 150000, "Pago virtual PSE", "C1"]])
        resultado = PseConciliador(pse, cruces).procesar()
        rows = [item["row"] for item in resultado["dataset"]]
        self.assertNotIn(2, rows)
        self.assertNotIn(3, rows)
        self.assertEqual(_states(resultado)[4], "Conciliado 1:1")

    def test_valor_negativo_usa_magnitud_exacta(self) -> None:
        pse = _pse([["2026-04-15", -150000, "Pago PSE", "P1"]])
        cruces = _cruces([["2026-04-15", 150000, "Pago virtual PSE", "C1"]])
        resultado = PseConciliador(pse, cruces).procesar()
        self.assertEqual(_states(resultado)[2], "Conciliado 1:1")


if __name__ == "__main__":
    unittest.main()
