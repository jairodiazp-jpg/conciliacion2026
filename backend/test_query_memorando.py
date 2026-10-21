from __future__ import annotations

import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from query_memorando_conciliador import conciliar_query_memorando


def workbook_bytes(rows, *, fill=False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Datos"
    sheet.append(["Cuenta", "Valor", "Fecha", "Documento", "SAP", "Estado"])
    for row in rows:
        sheet.append(row)
    if fill:
        sheet["A2"].fill = PatternFill(fill_type="solid", fgColor="FF00FF00")
        sheet.merge_cells("F4:G4")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class QueryMemorandoTests(unittest.TestCase):
    def run_case(self, query_rows, memo_rows):
        return conciliar_query_memorando(workbook_bytes(query_rows), workbook_bytes(memo_rows, fill=True))

    def test_exact_match_uses_memo_document_and_sap(self):
        result = self.run_case(
            [("001", "1000,00", "2026-04-01", "", "", "")],
            [("001", 1000, "2026-04-01", "DOC-1", "SAP-1", "")],
        )
        self.assertEqual(result["summary"]["coincidencias_exactas"], 1)
        self.assertEqual(result["results"][0]["memorando"]["documento"], "DOC-1")

    def test_different_account_does_not_match(self):
        result = self.run_case(
            [("001", 1000, "2026-04-01", "", "", "")],
            [("002", 1000, "2026-04-01", "DOC-1", "SAP-1", "")],
        )
        self.assertEqual(result["summary"]["coincidencias_exactas"], 0)

    def test_different_date_does_not_match(self):
        result = self.run_case(
            [("001", 1000, "2026-04-01", "", "", "")],
            [("001", 1000, "2026-04-02", "DOC-1", "SAP-1", "")],
        )
        self.assertEqual(result["summary"]["coincidencias_exactas"], 0)

    def test_different_cent_does_not_match(self):
        result = self.run_case(
            [("001", "1000,01", "2026-04-01", "", "", "")],
            [("001", "1000,00", "2026-04-01", "DOC-1", "SAP-1", "")],
        )
        self.assertEqual(result["summary"]["coincidencias_exactas"], 0)

    def test_missing_document_or_sap_is_not_confirmed(self):
        result = self.run_case(
            [("001", 1000, "2026-04-01", "", "", "")],
            [("001", 1000, "2026-04-01", "", "SAP-1", "")],
        )
        self.assertEqual(result["results"][0]["status"], "documento_o_sap_no_identificable")

    def test_ambiguous_candidates_are_not_consumed(self):
        result = self.run_case(
            [("001", 1000, "2026-04-01", "", "", "")],
            [
                ("001", 1000, "2026-04-01", "DOC-1", "SAP-1", ""),
                ("001", 1000, "2026-04-01", "DOC-2", "SAP-2", ""),
            ],
        )
        self.assertEqual(result["results"][0]["status"], "coincidencia_ambigua")

    def test_previously_conciliated_row_is_excluded(self):
        result = self.run_case(
            [("001", 1000, "2026-04-01", "", "", "")],
            [("001", 1000, "2026-04-01", "DOC-1", "SAP-1", "CRUCE 1")],
        )
        self.assertEqual(result["summary"]["sin_coincidencia"], 1)

    def test_memo_is_returned_unchanged(self):
        memo = workbook_bytes([("001", 1000, "2026-04-01", "DOC-1", "SAP-1", "")], fill=True)
        result = conciliar_query_memorando(workbook_bytes([("001", 1000, "2026-04-01", "", "", "")]), memo)
        returned = __import__("base64").b64decode(result["memorando_file"]["file"])
        before = load_workbook(BytesIO(memo))
        after = load_workbook(BytesIO(returned))
        self.assertEqual(before.sheetnames, after.sheetnames)
        self.assertEqual(before["Datos"]["A2"].fill.fgColor.rgb, after["Datos"]["A2"].fill.fgColor.rgb)
        self.assertEqual(list(before["Datos"].merged_cells.ranges), list(after["Datos"].merged_cells.ranges))


if __name__ == "__main__":
    unittest.main()
