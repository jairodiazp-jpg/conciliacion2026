from __future__ import annotations

import base64
import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from conciliador import GREEN_FILL, ORANGE_FILL, ConciliadorContable
from integrador import ProcesadorIntegrado
from test_pse_exactitud import _cruces, _pse


def _contable_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cuenta 1110"
    sheet.append(["Fecha", "Descripción", "Valor", "x", "y", "Aprobacion", "Comentario", "Observacion"])
    sheet.append(["sin registrar en libros"])
    sheet.append(["2026-04-10", "Libros", 1000, None, None, "APR-1", None, None])
    sheet.append(["sin registrar en el extracto"])
    sheet.append(["2026-04-10", "Extracto", 1000, None, None, "APR-1", None, None])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class TestRegresionExcel(unittest.TestCase):
    def test_conserva_hojas_encabezados_y_flujo_integrado(self) -> None:
        pse = _pse(
            [
                ["2026-04-15", 2500000, "Pago PSE 1", "P1"],
                ["2026-04-15", 1500000, "Pago PSE 2", "P2"],
                ["2026-04-15", 3000000, "Pago PSE 3", "P3"],
                ["2026-04-15", 3000000, "Pago PSE 4", "P4"],
            ]
        )
        cruces = _cruces([["2026-04-15", 10000000, "Pago virtual PSE", "C1"]])
        original_contable = load_workbook(BytesIO(_contable_bytes()))
        original_names = original_contable.sheetnames
        original_headers = [cell.value for cell in original_contable.active[1]]

        resultado = ProcesadorIntegrado(
            contable_bytes=_contable_bytes(),
            pse_bytes=pse,
            cruces_bytes=cruces,
        ).procesar()

        files = {item["name"]: item["file"] for item in resultado["files"]}
        self.assertIn("CONCILIACION_CONTABLE.xlsx", files)
        self.assertIn("PSE_CONCILIADO.xlsx", files)

        contable_wb = load_workbook(BytesIO(base64.b64decode(files["CONCILIACION_CONTABLE.xlsx"])))
        self.assertEqual(contable_wb.sheetnames, original_names)
        self.assertEqual([cell.value for cell in contable_wb.active[1]], original_headers)

        pse_wb = load_workbook(BytesIO(base64.b64decode(files["PSE_CONCILIADO.xlsx"])))
        self.assertEqual(pse_wb.sheetnames, ["Movimientos PSE"])
        self.assertEqual(pse_wb.active.cell(row=1, column=1).value, "Fecha")
        self.assertEqual(pse_wb.active.cell(row=2, column=2).fill.start_color.rgb, "FFFFC0CB")
        self.assertGreaterEqual(int(resultado["resumen"]["cruzados"]), 1)

    def test_conciliador_contable_conserva_colores_existentes(self) -> None:
        resultado = ConciliadorContable(_contable_bytes()).procesar()
        workbook = load_workbook(BytesIO(base64.b64decode(resultado["file"])))
        sheet = workbook.active
        self.assertEqual(sheet.title, "Cuenta 1110")
        libros = sheet.cell(row=3, column=3)
        extracto = sheet.cell(row=5, column=3)
        self.assertIn(
            libros.fill.start_color.rgb,
            {ORANGE_FILL.start_color.rgb, GREEN_FILL.start_color.rgb, "FFFFA500", "FF92D050"},
        )
        self.assertIn(
            extracto.fill.start_color.rgb,
            {ORANGE_FILL.start_color.rgb, GREEN_FILL.start_color.rgb, "FFFFA500", "FF92D050"},
        )


if __name__ == "__main__":
    unittest.main()
