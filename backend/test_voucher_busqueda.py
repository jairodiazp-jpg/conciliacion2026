from __future__ import annotations

import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from voucher_busqueda import search_voucher_documents


def excel_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Soportes"
    sheet.append(["Valor", "Fecha", "Punto de venta", "Referencia", "Cuenta"])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class VoucherSearchTests(unittest.TestCase):
    def test_exact_value_and_date_are_found(self):
        original = excel_bytes([(28144, "2026-07-03", "Tienda 1", "V-1", "1332")])
        result = search_voucher_documents(
            original,
            "memorando.xlsx",
            [("voucher.xlsx", original)],
            "28144.00",
            "2026-07-03",
        )
        self.assertEqual(result["summary"]["coincidencias"], 2)
        self.assertTrue(result["read_only"])

    def test_ambiguous_results_are_not_selected(self):
        original = excel_bytes([(100, "2026-01-01", "A", "1", "1"), (100, "2026-01-01", "B", "2", "2")])
        result = search_voucher_documents(original, "memo.xlsx", [], "100", "2026-01-01")
        self.assertTrue(all(item["status"] == "AMBIGUO" for item in result["results"]))

    def test_unsupported_source_is_reported(self):
        result = search_voucher_documents(excel_bytes([]), "memo.xlsx", [("scan.pdf", b"not indexed")], "1")
        self.assertEqual(result["summary"]["pendientes_revision"], 1)

    def test_sources_are_not_modified(self):
        original = excel_bytes([(100, "2026-01-01", "A", "1", "1")])
        before = load_workbook(BytesIO(original), data_only=False)
        search_voucher_documents(original, "memo.xlsx", [("voucher.xlsx", original)], "100")
        after = load_workbook(BytesIO(original), data_only=False)
        self.assertEqual(before["Soportes"]["A2"].value, after["Soportes"]["A2"].value)


if __name__ == "__main__":
    unittest.main()
