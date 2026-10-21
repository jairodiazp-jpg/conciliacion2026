from __future__ import annotations

import csv
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import PurePath
from typing import Any

from openpyxl import load_workbook


def _normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(char for char in text if not unicodedata.combining(char)).strip().lower()


def _decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"))
        except InvalidOperation:
            return None
    text = str(value).strip().replace("$", "").replace(" ", "")
    if not text:
        return None
    if "." in text and "," in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except (TypeError, ValueError):
            continue
    return None


def _field(header: Any) -> str | None:
    value = _normalize(header)
    if "valor" in value or "importe" in value or "monto" in value:
        return "value"
    if "fecha" in value:
        return "date"
    if "punto" in value or "establecimiento" in value or "comerc" in value:
        return "store"
    if "aprob" in value or "referencia" in value or "voucher" in value or "documento" in value:
        return "reference"
    if "cuenta" in value:
        return "account"
    return None


def _headers(values: list[Any]) -> dict[str, int]:
    found: dict[str, int] = {}
    for index, value in enumerate(values):
        field = _field(value)
        if field and field not in found:
            found[field] = index
    return found


def _read_spreadsheet(content: bytes, name: str) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    rows: list[dict[str, Any]] = []
    issues: list[dict[str, str]] = []
    suffix = PurePath(name).suffix.lower()
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
            parsed = list(csv.reader(StringIO(text)))
            sheet_name = "CSV"
            if parsed:
                headers = _headers(parsed[0])
                for number, values in enumerate(parsed[1:], 2):
                    rows.append({"values": values, "columns": headers, "sheet": sheet_name, "row": number})
            return rows, issues
        except (UnicodeDecodeError, csv.Error) as error:
            return [], [{"file": name, "status": "PENDIENTE DE REVISIÓN", "detail": f"CSV ilegible: {error}"}]
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as error:
        return [], [{"file": name, "status": "PENDIENTE DE REVISIÓN", "detail": f"Excel ilegible: {error}"}]
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        header_values = next(iterator, None)
        if header_values is None:
            continue
        columns = _headers(list(header_values))
        if "value" not in columns:
            issues.append({"file": name, "status": "PENDIENTE DE REVISIÓN", "detail": f"La hoja {sheet.title} no tiene columna de valor reconocible."})
            continue
        for number, values in enumerate(iterator, 2):
            rows.append({"values": list(values), "columns": columns, "sheet": sheet.title, "row": number})
    return rows, issues


def _record(row: dict[str, Any], name: str, source: str) -> dict[str, Any] | None:
    values = row["values"]
    columns = row["columns"]
    def get(field: str) -> Any:
        index = columns.get(field)
        return values[index] if index is not None and index < len(values) else None
    amount = _decimal(get("value"))
    if amount is None:
        return None
    parsed_date = _date(get("date"))
    return {
        "source": source,
        "file": name,
        "location": f"{row['sheet']} fila {row['row']}",
        "sheet": row["sheet"],
        "row": row["row"],
        "value": str(amount),
        "date": parsed_date.isoformat() if parsed_date else None,
        "store": str(get("store") or "").strip() or "PUNTO DE VENTA NO VERIFICADO",
        "reference": str(get("reference") or "").strip() or "NO VERIFICADO",
        "account": str(get("account") or "").strip() or "NO VERIFICADA",
        "evidence": " | ".join(str(item) for item in values if item not in (None, ""))[:1000],
        "ocr_confidence": None,
    }


def search_voucher_documents(
    memorando_content: bytes,
    memorando_name: str,
    voucher_files: list[tuple[str, bytes]],
    value: str,
    exact_date: str | None = None,
    store: str | None = None,
    reference: str | None = None,
    account: str | None = None,
) -> dict[str, Any]:
    query_value = _decimal(value)
    if query_value is None:
        raise ValueError("El valor exacto no es válido.")
    query_date = _date(exact_date) if exact_date else None
    memo_rows, memo_issues = _read_spreadsheet(memorando_content, memorando_name)
    sources = [("memorando", memorando_name, memo_rows)]
    issues = list(memo_issues)
    for name, content in voucher_files:
        suffix = PurePath(name).suffix.lower()
        if suffix not in {".xlsx", ".xlsm", ".csv"}:
            issues.append({"file": name, "status": "PENDIENTE DE REVISIÓN", "detail": "Formato no indexado: se requiere un extractor autorizado para PDF o imagen."})
            continue
        rows, file_issues = _read_spreadsheet(content, name)
        issues.extend(file_issues)
        sources.append(("voucher", name, rows))
    results: list[dict[str, Any]] = []
    for source, name, rows in sources:
        for row in rows:
            item = _record(row, name, source)
            if item is None or Decimal(item["value"]) != query_value:
                continue
            if query_date and item["date"] != query_date.isoformat():
                continue
            if store and _normalize(store) not in _normalize(item["store"]):
                continue
            if reference and _normalize(reference) not in _normalize(item["reference"]):
                continue
            if account and _normalize(account) not in _normalize(item["account"]):
                continue
            results.append(item)
    for item in results:
        item["status"] = "ENCONTRADO"
        item["observation"] = "Coincidencia exacta de valor; fecha y filtros aplicados cuando fueron informados."
    if len(results) > 1:
        for item in results:
            item["status"] = "AMBIGUO"
            item["observation"] = "Hay varios documentos con el mismo valor y filtros. Revise todas las evidencias."
    return {
        "summary": {
            "documentos_procesados": len(sources),
            "coincidencias": len(results),
            "pendientes_revision": len(issues),
        },
        "query": {"value": str(query_value), "date": query_date.isoformat() if query_date else None},
        "results": results,
        "issues": issues,
        "read_only": True,
        "warning": "Consulta de solo lectura. Los documentos originales no se escriben ni se guardan.",
    }
